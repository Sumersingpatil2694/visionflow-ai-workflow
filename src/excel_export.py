"""
Excel exporter — 5 sheets as per plan:
  1) Nodes
  2) Connections
  3) Statistics
  4) Summary
  5) Warnings / Confidence
"""
from __future__ import annotations

from typing import Any, Dict, List
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _style_header(ws, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER


def _autosize(ws) -> None:
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 8
        for cell in col:
            v = "" if cell.value is None else str(cell.value)
            if len(v) > max_len:
                max_len = min(60, len(v))
        ws.column_dimensions[col_letter].width = max_len + 2


def export_to_bytes(result: Dict[str, Any]) -> bytes:
    """Build the .xlsx in-memory and return bytes (for Streamlit download_button)."""
    wb = Workbook()

    # ---- Sheet 1: Nodes
    ws1 = wb.active
    ws1.title = "Nodes"
    ws1.append(["Node ID", "Label", "Shape", "X1", "Y1", "X2", "Y2", "Center X", "Center Y", "Confidence", "Source"])
    for n in result.get("nodes", []):
        x1, y1, x2, y2 = n["bbox"]
        cx, cy = n["center"]
        ws1.append([
            n["id"], n.get("label", ""), n.get("shape", ""),
            x1, y1, x2, y2, cx, cy,
            n.get("confidence", 0.0), n.get("source", ""),
        ])
    _style_header(ws1, 11); _autosize(ws1)

    # ---- Sheet 2: Connections
    ws2 = wb.create_sheet("Connections")
    ws2.append(["Edge ID", "Source ID", "Source Label", "Target ID", "Target Label", "Direction", "Confidence", "Arrowhead Detected"])
    id2label = {n["id"]: n.get("label", "") for n in result.get("nodes", [])}
    for e in result.get("edges", []):
        ws2.append([
            e["id"],
            e["source"], id2label.get(e["source"], ""),
            e["target"], id2label.get(e["target"], ""),
            e.get("direction", ""), e.get("confidence", 0.0),
            "Yes" if e.get("has_arrowhead") else "No",
        ])
    _style_header(ws2, 8); _autosize(ws2)

    # ---- Sheet 3: Statistics
    ws3 = wb.create_sheet("Statistics")
    ws3.append(["Metric", "Value"])
    stats = result.get("statistics", {})
    for k, v in stats.items():
        ws3.append([k.replace("_", " ").title(), v])
    _style_header(ws3, 2); _autosize(ws3)

    # ---- Sheet 4: Summary
    ws4 = wb.create_sheet("Summary")
    ws4.append(["Field", "Value"])
    ws4.append(["Summary", result.get("summary", "")])
    starts = ", ".join(result.get("start_states", []))
    ends = ", ".join(result.get("end_states", []))
    ws4.append(["Start States", starts])
    ws4.append(["End States", ends])
    for i, obs in enumerate(result.get("observations", []), start=1):
        ws4.append([f"Observation {i}", obs])
    _style_header(ws4, 2); _autosize(ws4)
    # allow wrap for summary
    ws4.column_dimensions["B"].width = 80
    for row in ws4.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    # ---- Sheet 5: Warnings & Confidence
    ws5 = wb.create_sheet("Warnings")
    ws5.append(["Type", "Message", "Severity"])
    for w in result.get("uncertainties", []):
        ws5.append(["Uncertainty", w, "Info"])
    # low-confidence nodes
    for n in result.get("nodes", []):
        if n.get("confidence", 1.0) < 0.45:
            ws5.append(["Low-Confidence Node", f"{n['id']} - '{n.get('label','')}' (conf={n.get('confidence')})", "Warning"])
    # low-confidence edges
    for e in result.get("edges", []):
        if e.get("confidence", 1.0) < 0.40:
            ws5.append(["Low-Confidence Edge", f"{e['id']} {e['source']}->{e['target']} (conf={e.get('confidence')})", "Warning"])
    _style_header(ws5, 3); _autosize(ws5)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
